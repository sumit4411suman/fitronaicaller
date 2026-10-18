"""Blueprint features: roles/JWT, workers, campaign engine, follow-up scheduler, reports, calendar, sentiment,
call summary, streaming LLM, real-time transcription gates, PDF/DOCX knowledge import, encryption at rest."""

import io
import json
from datetime import UTC, datetime, timedelta

import numpy as np
import pytest

from tests.conftest import LOCAL_TOKEN, make_contact, open_hours

pytestmark = pytest.mark.automated


def _user(client, username, role, pin="482913"):
    r = client.post("/api/v1/users", json={"username": username, "display_name": username.title(), "role": role, "pin": pin})
    assert r.status_code == 201, r.text
    tok = client.post("/api/v1/auth/login", json={"username": username, "pin": pin}).json()["token"]
    return r.json()["id"], {"Authorization": f"Bearer {tok}"}


# ----------------------------------------------------------------------------- roles & JWT
def test_roles_and_permissions(client):
    _, viewer = _user(client, "vik", "viewer")
    _, operator = _user(client, "ops", "operator")
    _, manager = _user(client, "mgr", "manager")
    assert client.get("/api/v1/auth/whoami", headers=viewer).json()["role"] == "viewer"
    assert client.get("/api/v1/contacts", headers=viewer).status_code == 200
    assert client.post("/api/v1/contacts", headers=viewer, json={"business_name": "X", "phone": "9876543210"}).status_code == 403
    assert client.post("/api/v1/contacts", headers=operator, json={"business_name": "X", "phone": "9876543210"}).status_code == 201
    assert client.patch("/api/v1/settings", headers=operator, json={"values": {"calling.max_calls_per_day": 5}}).status_code == 403
    assert client.patch("/api/v1/settings", headers=manager, json={"values": {"calling.max_calls_per_day": 5}}).status_code == 403
    assert client.get("/api/v1/users", headers=manager).status_code == 403
    assert client.get("/api/v1/audit", headers=operator).status_code == 403
    assert client.get("/api/v1/audit", headers=manager).status_code == 200
    doc = client.post("/api/v1/knowledge", headers=manager, json={"title": "FAQ", "body": "FITRON is gym accounting software.", "category": "faq"}).json()
    assert client.post(f"/api/v1/knowledge/{doc['id']}/approve", headers=operator, json={}).status_code == 403
    assert client.post(f"/api/v1/knowledge/{doc['id']}/approve", headers=manager, json={}).status_code == 200
    # emergency stop: any role may activate, only calls:write may clear
    assert client.post("/api/v1/call-jobs/emergency-stop", headers=viewer, json={"reason": "viewer saw a problem"}).json()["active"]
    assert client.post("/api/v1/call-jobs/emergency-stop", headers=viewer, json={"reason": "x", "activate": False}).status_code == 403
    assert client.post("/api/v1/call-jobs/emergency-stop", headers=operator, json={"reason": "ok", "activate": False}).json()["active"] is False
    audit = client.get("/api/v1/audit", params={"action": "emergency_stop"}).json()
    assert any(a["actor"] == "vik" for a in audit)


def test_jwt_revocation_lockout_and_weak_pins(client):
    uid, h = _user(client, "asha", "operator")
    assert client.patch(f"/api/v1/users/{uid}", json={"role": "viewer"}).status_code == 200
    assert client.get("/api/v1/contacts", headers=h).status_code == 401  # role change revokes old JWT
    assert client.post("/api/v1/users", json={"username": "weak", "display_name": "W", "role": "viewer", "pin": "111111"}).status_code == 400
    assert client.post("/api/v1/users", json={"username": "seq", "display_name": "S", "role": "viewer", "pin": "123456"}).status_code == 400
    for _ in range(5):
        assert client.post("/api/v1/auth/login", json={"username": "asha", "pin": "000999"}).status_code == 401
    assert client.post("/api/v1/auth/login", json={"username": "asha", "pin": "482913"}).status_code == 423
    client.post(f"/api/v1/users/{uid}/reset-pin", json={"pin": "739205"})
    assert client.post("/api/v1/auth/login", json={"username": "asha", "pin": "739205"}).status_code == 200
    client.patch(f"/api/v1/users/{uid}", json={"active": False})
    assert client.post("/api/v1/auth/login", json={"username": "asha", "pin": "739205"}).status_code == 401
    forged = "eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiJ4Iiwicm9sZSI6ImFkbWluIn0.bad"
    assert client.get("/api/v1/contacts", headers={"Authorization": f"Bearer {forged}"}).status_code == 401


def test_jwt_on_desktop_websocket(client):
    _, h = _user(client, "wsuser", "viewer")
    tok = h["Authorization"].split()[1]
    with client.websocket_connect(f"/api/v1/ws/desktop?token={tok}") as ws:
        client.post("/api/v1/call-jobs/emergency-stop", json={"reason": "t"})
        assert json.loads(ws.receive_text())["type"] == "safety.emergency_stop"


# ----------------------------------------------------------------------------- workers
def test_call_queue_worker_queues_due_jobs(client):
    open_hours(client)
    c = make_contact(client)
    job = client.post("/api/v1/call-jobs", json={"contact_id": c["id"], "scheduled_for": (datetime.now(UTC) - timedelta(minutes=1)).isoformat()}).json()
    assert job["status"] == "pending"
    with client.websocket_connect(f"/api/v1/ws/desktop?token={LOCAL_TOKEN}") as ws:
        r = client.post("/api/v1/system/workers/call_queue/run").json()
        assert r["result"]["queued"] == 1
        ev = json.loads(ws.receive_text())
        assert ev["type"] == "event" and ev["payload"]["name"] == "queue.ready"
    assert client.get(f"/api/v1/call-jobs/{job['id']}").json()["status"] == "queued"
    st = client.get("/api/v1/system/workers").json()
    assert set(st["periodic"]) == {"call_queue", "campaign_engine", "followup_scheduler", "report"} and st["ai"]["name"] == "ai"


def test_campaign_start_engine_and_progress(client):
    open_hours(client)
    a = make_contact(client)
    b = make_contact(client, business_name="No Consent Gym", phone="9123456780", consent_status="unknown")
    camp = client.post("/api/v1/campaigns", json={"name": "Bokaro gyms", "auto_followup": True, "followup_after_days": 3}).json()
    client.post(f"/api/v1/campaigns/{camp['id']}/contacts", json={"ids": [a["id"], b["id"]]})
    r = client.post(f"/api/v1/campaigns/{camp['id']}/start").json()
    assert r["campaign"]["status"] == "active" and r["engine"]["created"] == 1
    jobs = client.get("/api/v1/call-jobs", params={"status": "pending"}).json()["items"]
    assert [j["contact_id"] for j in jobs] == [a["id"]]
    r2 = client.post("/api/v1/system/workers/campaign_engine/run").json()["result"]
    assert r2["excluded"] == 1  # no consent -> excluded, never dialled
    stats = client.get(f"/api/v1/campaigns/{camp['id']}/stats").json()
    assert stats["excluded"] == 1 and "progress_percent" in stats
    assert client.post(f"/api/v1/campaigns/{camp['id']}/pause").json()["status"] == "paused"


def test_followup_scheduler_reminds_creates_jobs_and_auto_followups(client):
    open_hours(client)
    c = make_contact(client)
    due = datetime.now(UTC) + timedelta(minutes=5)
    fu = client.post("/api/v1/followups", json={"contact_id": c["id"], "due_at": due.isoformat(), "kind": "call", "note": "call back"}).json()
    with client.websocket_connect(f"/api/v1/ws/desktop?token={LOCAL_TOKEN}") as ws:
        r = client.post("/api/v1/system/workers/followup_scheduler/run").json()["result"]
        assert r["reminded"] == 1 and r["jobs_created"] == 1
        ev = json.loads(ws.receive_text())
        assert ev["payload"]["name"] == "followup.due" and ev["payload"]["data"]["followup_id"] == fu["id"]
    # second run does not remind twice
    assert client.post("/api/v1/system/workers/followup_scheduler/run").json()["result"]["reminded"] == 0
    # automatic follow-up after a no-answer call
    c2 = make_contact(client, business_name="Iron House", phone="9123456780")
    job = client.post("/api/v1/call-jobs", json={"contact_id": c2["id"]}).json()
    call_id = client.post(f"/api/v1/call-jobs/{job['id']}/dispatch").json()["call_id"]
    client.post(f"/api/v1/calls/{call_id}/no-answer")
    r = client.post("/api/v1/system/workers/followup_scheduler/run").json()["result"]
    assert r["auto_created"] == 1
    fus = client.get("/api/v1/followups", params={"contact_id": c2["id"]}).json()
    assert len(fus) == 1 and "no_answer" in fus[0]["note"]


def test_reports_generate_download_and_path_safety(client):
    make_contact(client)
    info = client.post("/api/v1/reports/generated").json()
    assert set(info["files"]) == {"summary.json", "calls.csv", "report.html"}
    listed = client.get("/api/v1/reports/generated").json()
    assert listed[0]["date"] == info["date"]
    html = client.get(f"/api/v1/reports/generated/{info['date']}/report.html")
    assert html.status_code == 200 and "SELL <span>AGENT</span>" in html.text
    assert client.get(f"/api/v1/reports/generated/{info['date']}/..%2F..%2Flocal_api_token").status_code == 404
    assert client.get("/api/v1/reports/generated/../x/summary.json").status_code == 404
    assert client.post("/api/v1/reports/generated", params={"date": "not-a-date"}).status_code == 400


def test_calendar_and_dashboard_comparisons(client):
    c = make_contact(client)
    due = datetime.now(UTC) + timedelta(days=1)
    client.post("/api/v1/followups", json={"contact_id": c["id"], "due_at": due.isoformat(), "kind": "demo"})
    cal = client.get("/api/v1/followups/calendar", params={"month": due.strftime("%Y-%m")}).json()
    assert cal["total"] >= 1 and any(item["kind"] == "demo" for items in cal["days"].values() for item in items)
    assert client.get("/api/v1/followups/calendar", params={"month": "2026-13"}).status_code == 400
    dash = client.get("/api/v1/reports/dashboard").json()
    assert dash["comparisons"]["contacts_added_week_vs_previous"] == {"current": 1, "previous": 0}


# ----------------------------------------------------------------------------- sentiment & summary
def test_sentiment_and_call_summary(client):
    open_hours(client)
    c = make_contact(client)
    job = client.post("/api/v1/call-jobs", json={"contact_id": c["id"]}).json()
    call_id = client.post(f"/api/v1/call-jobs/{job['id']}/dispatch").json()["call_id"]
    client.post(f"/api/v1/calls/{call_id}/connected")
    client.post(f"/api/v1/calls/{call_id}/step", json={"utterance": None, "language": "en", "use_llm": False})
    t = client.post(f"/api/v1/calls/{call_id}/step", json={"utterance": "yes sure, sounds great", "language": "en", "use_llm": False}).json()
    assert t["sentiment"] == "positive"
    client.post(f"/api/v1/calls/{call_id}/step", json={"utterance": "how much is the price?", "language": "en", "use_llm": False})
    client.post(f"/api/v1/calls/{call_id}/outcome", json={"outcome": "demo_scheduled", "followup_at": (datetime.now(UTC) + timedelta(days=2)).isoformat(), "followup_kind": "demo"})
    s = client.get(f"/api/v1/calls/{call_id}/summary").json()
    assert s["result"] == "demo_scheduled" and s["sentiment"] == "positive"
    assert "Asked about pricing" in s["key_points"] and s["next_action"]["kind"] == "demo" and s["knowledge_used"]


@pytest.mark.parametrize("text,label", [("bahut accha, demo bhejo", "positive"), ("नहीं चाहिए, बहुत महंगा है", "negative"), ("call mat karo", "negative"), ("I am not interested", "negative"), ("what time is it", "neutral")])
def test_sentiment_lexicon(text, label):
    from sellagent_ai.sentiment import analyze

    assert analyze(text).label == label


# ----------------------------------------------------------------------------- streaming LLM
class _FakeLLM:
    name = "fake"

    def available(self):
        return True, "ok"

    def list_models(self):
        return []

    def stream(self, messages, *, model, temperature=0.3, max_tokens=256, cancel=None, timeout=60.0):
        yield from ["Namaste", ", FITRON", " is ready."]


def test_chat_stream_sse(client, monkeypatch):
    import sellagent_backend.api.v1.ai as ai_mod

    monkeypatch.setattr(ai_mod, "make_provider", lambda *a, **k: _FakeLLM())
    with client.stream("POST", "/api/v1/ai/chat-stream", json={"prompt": "hello", "language": "en"}) as r:
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/event-stream")
        body = "".join(r.iter_text())
    tokens = [json.loads(line[6:])["token"] for line in body.splitlines() if line.startswith("data: ") and '"token"' in line]
    assert "".join(tokens) == "Namaste, FITRON is ready."
    assert "event: done" in body and '"status": "done"' in body


def test_chat_stream_reports_unavailable_provider(client):
    with client.stream("POST", "/api/v1/ai/chat-stream", json={"prompt": "hello"}) as r:
        body = "".join(r.iter_text())
    assert "event: error" in body and "event: done" in body


def test_qwen25_requires_licence_ack(client):
    cat = client.get("/api/v1/ai/catalog").json()["candidates"]
    q = [c for c in cat if c["name"] == "qwen2.5:3b"][0]
    assert "NON-COMMERCIAL" in q["license"]
    r = client.post("/api/v1/ai/models/pull", json={"model": "qwen2.5:3b", "acknowledge_download": True})
    assert r.status_code == 400 and r.json()["code"] == "license_ack_required"


# ----------------------------------------------------------------------------- real-time transcription
class _FakeSTT:
    model_size = "base"

    def available(self):
        return True, "fake"

    def transcribe(self, audio, *, language=None, allow_download=False):
        from sellagent_speech.stt import STTResult

        return STTResult(text="haan ji, bahut accha", language="hi", language_probability=0.9, duration_s=len(audio) / 16000, latency_ms=5.0, model="fake", segments=[])


def _live_call(client):
    open_hours(client)
    c = make_contact(client)
    job = client.post("/api/v1/call-jobs", json={"contact_id": c["id"]}).json()
    call_id = client.post(f"/api/v1/call-jobs/{job['id']}/dispatch").json()["call_id"]
    client.post(f"/api/v1/calls/{call_id}/connected")
    return call_id


def test_transcription_gates(client):
    call_id = _live_call(client)
    with client.websocket_connect(f"/api/v1/ws/transcribe?token={LOCAL_TOKEN}&call_id={call_id}") as ws:
        assert "disabled" in json.loads(ws.receive_text())["error"]
    client.patch("/api/v1/settings", json={"values": {"privacy.allow_live_transcription": True}})
    with client.websocket_connect(f"/api/v1/ws/transcribe?token={LOCAL_TOKEN}&call_id={call_id}") as ws:
        assert "consent" in json.loads(ws.receive_text())["error"]
    assert client.post(f"/api/v1/calls/{call_id}/transcription-consent", json={"prospect_consented": True, "disclosure_text_read": False}).status_code == 400


def test_transcription_streams_text_without_storing_audio(client, settings):
    call_id = _live_call(client)
    client.patch("/api/v1/settings", json={"values": {"privacy.allow_live_transcription": True}})
    client.post(f"/api/v1/calls/{call_id}/transcription-consent", json={"prospect_consented": True, "disclosure_text_read": True})
    client.app.state.stt_engine = _FakeSTT()
    sr = 16000
    tone = (0.3 * np.sin(2 * np.pi * 220 * np.arange(int(sr * 1.0)) / sr)).astype(np.float32)
    pcm = np.concatenate([np.zeros(sr // 2, np.float32), tone, np.zeros(sr, np.float32)])
    data = (pcm * 32767).astype("<i2").tobytes()
    with client.websocket_connect(f"/api/v1/ws/transcribe?token={LOCAL_TOKEN}&call_id={call_id}&language=hi") as ws:
        assert json.loads(ws.receive_text())["type"] == "ready"
        for i in range(0, len(data), 3200):
            ws.send_bytes(data[i : i + 3200])
        msgs = []
        while True:
            m = json.loads(ws.receive_text())
            msgs.append(m)
            if m["type"] == "final":
                break
        ws.send_text(json.dumps({"type": "stop"}))
    final = msgs[-1]
    assert final["text"] == "haan ji, bahut accha" and final["sentiment"] == "positive"
    tr = client.get(f"/api/v1/calls/{call_id}").json()["transcript"]
    assert any(t["text"] == "haan ji, bahut accha" and t["speaker"] == "prospect" for t in tr)
    assert not any(p.suffix in (".wav", ".pcm", ".raw") for p in settings.data_dir.rglob("*"))


# ----------------------------------------------------------------------------- PDF / DOCX knowledge import
def _pdf(text: str) -> bytes:
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{o}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_pdf_and_docx_import(client):
    from docx import Document

    r = client.post("/api/v1/knowledge/import", files={"file": ("brochure.pdf", _pdf("FITRON monthly billing for gyms"), "application/pdf")}, params={"category": "faq"})
    assert r.status_code == 201, r.text
    assert "FITRON monthly billing" in r.json()["body"] and r.json()["status"] == "pending_review" and not r.json()["trusted"]
    d = Document()
    d.add_paragraph("FITRON objection: we already use Excel.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Plan", "Professional"
    buf = io.BytesIO()
    d.save(buf)
    r2 = client.post("/api/v1/knowledge/import", files={"file": ("objections.docx", buf.getvalue(), "application/octet-stream")}, params={"category": "objection"})
    assert r2.status_code == 201 and "Excel" in r2.json()["body"] and "Professional" in r2.json()["body"]
    assert client.post("/api/v1/knowledge/import", files={"file": ("bad.pdf", b"not a pdf", "application/pdf")}).status_code == 400
    assert client.post("/api/v1/knowledge/import", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 400


# ----------------------------------------------------------------------------- encryption at rest
def test_database_file_is_encrypted(client, settings):
    from sellagent_backend.db import crypto

    make_contact(client, business_name="Secret Gym Name")
    h = client.get("/api/v1/health").json()["database"]
    if not crypto.available():
        assert h["encrypted"] is False
        pytest.skip("SQLCipher not installed")
    assert h["encrypted"] is True and h["encryption"].startswith("sqlcipher")
    client.app.state.engine.dispose()
    blob = settings.db_path.read_bytes() + (settings.db_path.with_name(settings.db_path.name + "-wal").read_bytes() if settings.db_path.with_name(settings.db_path.name + "-wal").exists() else b"")
    assert b"Secret Gym Name" not in blob and b"+919876543210" not in blob and not blob.startswith(crypto.PLAINTEXT_HEADER)
